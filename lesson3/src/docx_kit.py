# -*- coding: utf-8 -*-
"""打印资料的排版工具 —— 复用第二节课打印资料的样式表。"""

import copy
import os

from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "assets", "lesson2_template.docx")

FULL_W = 9866            # 正文可用宽度（twips）
BORDER = "C9D6E2"
BORDER_LIGHT = "E2E7EC"
FILL_INFO = "EAF4FF"
FILL_SOFT = "F9FBFD"
FILL_HEAD = "DCEEFF"
INK_MUTED = "5B6570"
ACCENT = "3D8DFF"


# ------------------------------------------------------------------ 基础
def load(header="IELTS READING 精讲第三课 • 学生打印材料"):
    """打开模板并清空正文，只保留样式、页眉页脚与页面设置。"""
    doc = Document(TEMPLATE)
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.runs:
                p.runs[0].text = header
                for extra in p.runs[1:]:
                    extra.text = ""
    return doc


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _borders(cell, color=BORDER, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:color"), color)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        b.append(e)
    tcPr.append(b)


def _set_widths(table, widths):
    """固定列宽：必须同时写 tblW / tblLayout / tblGrid / tcW，Word 与 LibreOffice 才都认。"""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr

    for tag, attrs in (("w:tblW", {"w:w": str(sum(widths)), "w:type": "dxa"}),
                       ("w:tblLayout", {"w:type": "fixed"})):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        tblPr.append(el)

    for old in tbl.findall(qn("w:tblGrid")):
        tbl.remove(old)
    grid_el = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid_el.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid_el)

    for row in table.rows:
        # 行内容不跨页断开，打印时更易读
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for cell, w in zip(row.cells, widths):
            cell.width = Twips(w)


def _cell_text(cell, runs, size=9, align=None, space_after=2):
    """runs: str 或 [(text, bold, color), ...]"""
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(space_after)
    items = [(runs, False, None)] if isinstance(runs, str) else runs
    for text, bold, color in items:
        for j, seg in enumerate(str(text).split("\n")):
            if j:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(space_after)
                if align:
                    p.alignment = align
            r = p.add_run(seg)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Arial Unicode MS"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


# ------------------------------------------------------------------ 段落
def para(doc, text, style=None, size=None, bold=False, color=None,
         space_after=None, italic=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if italic:
            r.font.italic = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r.font.name = "Arial Unicode MS"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    return p


def rich(doc, segs, style=None, space_after=None):
    """segs: [(text, bold, color, size), ...]"""
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    for text, bold, color, size in segs:
        r = p.add_run(text)
        r.font.bold = bold
        r.font.size = Pt(size or 10)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r.font.name = "Arial Unicode MS"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    return p


def note_box(doc, label, text, fill=FILL_INFO, size=9):
    """整幅提示框（对应第二节课的“使用建议 / 建议”框）。"""
    t = doc.add_table(rows=1, cols=1)
    _set_widths(t, [FULL_W])
    c = t.cell(0, 0)
    _shade(c, fill)
    _borders(c, BORDER)
    segs = []
    if label:
        segs.append((f"{label}　", True, None))
    segs.append((text, False, None))
    _cell_text(c, segs, size=size)
    para(doc, "", space_after=3)
    return t


def image(doc, path, width_in=6.3):
    from docx.shared import Inches
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width_in))
    return p


def grid(doc, header, rows, widths, head_fill=FILL_HEAD, size=9,
         body_fill=None, first_bold=False, align=None):
    """带表头的表格。header 为 None 时不画表头。"""
    n = len(widths)
    t = doc.add_table(rows=(1 if header else 0) + len(rows), cols=n)
    _set_widths(t, widths)
    ri = 0
    if header:
        for j, h in enumerate(header):
            c = t.cell(0, j)
            _shade(c, head_fill)
            _borders(c, BORDER_LIGHT, 2)
            _cell_text(c, [(h, True, None)], size=size)
        ri = 1
    for row in rows:
        for j, val in enumerate(row):
            c = t.cell(ri, j)
            _borders(c, BORDER_LIGHT, 2)
            if body_fill:
                _shade(c, body_fill)
            bold = first_bold and j == 0
            _cell_text(c, [(val, bold, None)], size=size, align=align)
        ri += 1
    para(doc, "", space_after=3)
    return t


def blank_analysis_table(doc, sentences):
    """课堂用“长难句拆解区”：给出原句，主干/逻辑留白。"""
    rows = [(s, "", "") for s in sentences]
    return grid(doc, ("原句", "主干 / 从句类型", "逻辑 / 简化表达"), rows,
                [5100, 2400, 2366], size=8.5)


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


CENTER = WD_ALIGN_PARAGRAPH.CENTER
