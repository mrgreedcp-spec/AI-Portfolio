# -*- coding: utf-8 -*-
"""Slide engine for the TOEFL Academic Discussion Lesson 2 classroom deck.

Layout rules enforced here:
  * 16:9 projection canvas (13.333in x 7.5in)
  * every substantive body run is >= 19pt (spec asks for >= 18.5pt)
  * one running footer with part label + slide number
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

SW, SH = 13.333, 7.5
ML, MR = 0.68, 0.68
CW = SW - ML - MR

INK        = RGBColor(0x17, 0x2A, 0x31)
TEAL       = RGBColor(0x0E, 0x53, 0x62)
TEAL_MID   = RGBColor(0x2A, 0x7E, 0x8E)
TEAL_PALE  = RGBColor(0xE4, 0xEF, 0xF1)
TEAL_TINT  = RGBColor(0xF3, 0xF8, 0xF9)
ACCENT     = RGBColor(0xC9, 0x55, 0x27)
ACCENT_PALE= RGBColor(0xFB, 0xEC, 0xE3)
GOLD       = RGBColor(0xB0, 0x83, 0x00)
GOLD_PALE  = RGBColor(0xFA, 0xF2, 0xDC)
GREEN      = RGBColor(0x2E, 0x6B, 0x4F)
GREEN_PALE = RGBColor(0xE7, 0xF1, 0xEB)
GREY       = RGBColor(0x62, 0x73, 0x78)
LINE       = RGBColor(0xCF, 0xDD, 0xE1)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
PAPER      = RGBColor(0xFC, 0xFD, 0xFD)

LATIN = "Segoe UI"
EA    = "Microsoft YaHei"

BODY = 19          # minimum body size
LABEL = 19         # small labels (not body copy)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.part = ""
        self.n = 0

    # ---------- primitives ----------
    def slide(self, footer=True, bg=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        if bg is not None:
            self.rect(s, 0, 0, SW, SH, fill=bg, line=None)
        if footer:
            self._footer(s)
        return s

    def _footer(self, s):
        self.line(s, ML, 6.93, ML + CW, 6.93, LINE, 0.75)
        t = self.tb(s, ML, 6.99, CW * 0.7, 0.34)
        self.p(t, self.part, size=12, color=GREY, ea_bold=False)
        t2 = self.tb(s, ML + CW * 0.7, 6.99, CW * 0.3, 0.34)
        self.p(t2, f"{self.n}", size=12, color=GREY, align=PP_ALIGN.RIGHT)

    def rect(self, s, x, y, w, h, fill=None, line=None, lw=1.0, rounded=False, adj=0.06):
        shp = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            try:
                shp.adjustments[0] = adj
            except Exception:
                pass
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(lw)
        shp.shadow.inherit = False
        shp.text_frame.word_wrap = True
        return shp

    def line(self, s, x1, y1, x2, y2, color=LINE, w=1.0):
        c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(w)
        return c

    def tb(self, s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        tf._first = True
        return tf

    def p(self, tf, text="", size=BODY, bold=False, color=INK, align=PP_ALIGN.LEFT,
          space_before=0, space_after=4, line_spacing=1.16, italic=False,
          latin=LATIN, ea=EA, ea_bold=None, bullet=None, indent=0.0):
        if getattr(tf, "_first", False):
            para = tf.paragraphs[0]
            tf._first = False
        else:
            para = tf.add_paragraph()
        para.alignment = align
        para.space_before = Pt(space_before)
        para.space_after = Pt(space_after)
        para.line_spacing = line_spacing
        if indent:
            para.paragraph_format.left_indent = Inches(indent) if hasattr(para, "paragraph_format") else None
            pPr = para._p.get_or_add_pPr()
            pPr.set('marL', str(Emu(Inches(indent)).emu if False else int(indent * 914400)))
            pPr.set('indent', str(-int(0.22 * 914400)))
        if bullet is not None:
            text = f"{bullet}  {text}"
        self.run(para, text, size, bold, color, italic, latin, ea)
        return para

    def run(self, para, text, size=BODY, bold=False, color=INK, italic=False,
            latin=LATIN, ea=EA, underline=False):
        r = para.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.underline = underline
        f.color.rgb = color
        f.name = latin
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", ea if tag == "a:ea" else latin)
        return r

    def rich(self, tf, chunks, size=BODY, align=PP_ALIGN.LEFT, space_after=4,
             line_spacing=1.16, color=INK):
        """chunks: list of (text, dict-of-overrides)"""
        if getattr(tf, "_first", False):
            para = tf.paragraphs[0]
            tf._first = False
        else:
            para = tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        para.line_spacing = line_spacing
        for text, ov in chunks:
            self.run(para, text, ov.get("size", size), ov.get("bold", False),
                     ov.get("color", color), ov.get("italic", False),
                     ov.get("latin", LATIN), ov.get("ea", EA), ov.get("underline", False))
        return para

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    # ---------- composed blocks ----------
    def header(self, s, title, kicker=None, color=TEAL, rule=True, y=0.44):
        if kicker:
            t = self.tb(s, ML, y, CW, 0.32)
            self.p(t, kicker, size=LABEL, bold=True, color=ACCENT, space_after=0)
            ty = y + 0.36
        else:
            ty = y + 0.02
        t = self.tb(s, ML, ty, CW, 0.72)
        self.p(t, title, size=31, bold=True, color=color, space_after=0, line_spacing=1.05)
        if rule:
            self.line(s, ML, ty + 0.78, ML + 2.1, ty + 0.78, ACCENT, 3.0)
        return ty + 0.98

    def card(self, s, x, y, w, h, title=None, fill=TEAL_TINT, line=LINE,
             title_color=TEAL, title_size=20, pad=0.28):
        self.rect(s, x, y, w, h, fill=fill, line=line, rounded=True)
        cy = y + pad * 0.72
        if title:
            t = self.tb(s, x + pad, cy, w - 2 * pad, 0.34)
            self.p(t, title, size=title_size, bold=True, color=title_color, space_after=0)
            cy += 0.44
        return cy

    def tag(self, s, x, y, text, fill=ACCENT, color=WHITE, w=None, size=17, h=0.38):
        w = w or (0.26 + 0.132 * len(text))
        self.rect(s, x, y, w, h, fill=fill, line=None, rounded=True, adj=0.4)
        t = self.tb(s, x, y + 0.045, w, h, anchor=MSO_ANCHOR.TOP)
        self.p(t, text, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, space_after=0)
        return w

    def writing_lines(self, s, x, y, w, count, gap=0.44, color=LINE):
        for i in range(count):
            self.line(s, x, y + i * gap, x + w, y + i * gap, color, 1.0)
        return y + (count - 1) * gap

    def section(self, title, en, minutes, goal, steps=None, num=""):
        s = self.slide(bg=TEAL)
        self.rect(s, 0, 0, 0.34, SH, fill=ACCENT, line=None)
        t = self.tb(s, 1.5, 1.62, CW - 1.0, 0.5)
        self.p(t, num, size=46, bold=True, color=RGBColor(0x5C, 0x94, 0xA0), space_after=2)
        t = self.tb(s, 1.5, 2.62, CW - 1.0, 0.8)
        self.p(t, title, size=42, bold=True, color=WHITE, space_after=4)
        t = self.tb(s, 1.5, 3.5, CW - 1.0, 0.45)
        self.p(t, en, size=22, color=RGBColor(0x9E, 0xC6, 0xCE), space_after=0)
        self.rect(s, 1.5, 4.25, 0.06, 0.7, fill=ACCENT, line=None)
        t = self.tb(s, 1.78, 4.25, CW - 2.0, 0.8)
        self.p(t, f"本环节目标：{goal}", size=BODY, color=WHITE, space_after=0)
        self.tag(s, 1.5, 5.35, f"{minutes} 分钟", fill=ACCENT, size=17, h=0.4)
        if steps:
            t = self.tb(s, 3.5, 5.40, CW - 3.2, 0.5)
            self.p(t, " · ".join(steps), size=19, color=RGBColor(0x9E, 0xC6, 0xCE), space_after=0)
        return s

    def save(self, path):
        self.prs.save(path)
        return path


# --------------------------------------------------------------------------
# text metrics (used by layouts that must size themselves to their content)
# --------------------------------------------------------------------------
from PIL import ImageFont as _IF
_FONT_FILE = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
_fcache = {}


def _f(px):
    px = max(6, int(round(px)))
    if px not in _fcache:
        _fcache[px] = _IF.truetype(_FONT_FILE, px)
    return _fcache[px]


def est_lines(text, size_pt, width_in):
    if not text:
        return 1
    f = _f(size_pt / 72.0 * 96.0)
    maxw = width_in * 96.0
    toks, buf = [], ""
    for ch in text:
        if ord(ch) > 0x2E80:
            if buf:
                toks.append(buf); buf = ""
            toks.append(ch)
        elif ch == " ":
            buf += ch; toks.append(buf); buf = ""
        else:
            buf += ch
    if buf:
        toks.append(buf)
    lines, cur = 1, ""
    for t in toks:
        trial = cur + t
        if f.getlength(trial.rstrip()) > maxw and cur.strip():
            lines += 1
            cur = t.lstrip()
        else:
            cur = trial
    return lines


def est_h(text, size_pt, width_in, line_spacing=1.16):
    """Rendered height in inches."""
    return est_lines(text, size_pt, width_in) * size_pt * line_spacing * 1.03 / 72.0
