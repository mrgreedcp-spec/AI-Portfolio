# -*- coding: utf-8 -*-
"""Layout QA for the generated deck.

PowerPoint does not clip overflowing text, it spills it below the box. So the
real failure modes are (a) text running past the footer rule / off the slide and
(b) two text blocks landing on top of each other. This script measures every
text frame with a real CJK-capable font and reports both.
"""
import sys, math
from pptx import Presentation
from pptx.util import Pt
from PIL import ImageFont

FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
EMU_IN = 914400.0
FOOT_Y = 6.90          # footer rule
SAFE_BOTTOM = 7.30     # hard slide edge tolerance
_cache = {}


def font(px):
    px = max(6, int(round(px)))
    if px not in _cache:
        _cache[px] = ImageFont.truetype(FONT_PATH, px)
    return _cache[px]


def measure(text, size_pt, width_in):
    """Return number of wrapped lines for text at size_pt inside width_in."""
    if not text:
        return 1
    px_per_in = 96.0
    f = font(size_pt / 72.0 * px_per_in)
    maxw = width_in * px_per_in
    lines, cur = 1, ""
    # greedy wrap: break on spaces for latin, on any char for CJK
    tokens, buf = [], ""
    for ch in text:
        if ord(ch) > 0x2E80:                     # CJK / fullwidth
            if buf:
                tokens.append(buf); buf = ""
            tokens.append(ch)
        elif ch == " ":
            buf += ch
            tokens.append(buf); buf = ""
        else:
            buf += ch
    if buf:
        tokens.append(buf)
    for tok in tokens:
        trial = cur + tok
        if f.getlength(trial.rstrip()) > maxw and cur.strip():
            lines += 1
            cur = tok.lstrip()
        else:
            cur = trial
    return lines


def frame_height(tf, width_in):
    """Estimated rendered height in inches, plus max font size used."""
    total, maxsize = 0.0, 0.0
    for para in tf.paragraphs:
        txt = "".join(r.text for r in para.runs)
        sizes = [r.font.size.pt for r in para.runs if r.font.size]
        sz = max(sizes) if sizes else 18.0
        maxsize = max(maxsize, sz)
        ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.16
        sa = para.space_after.pt if para.space_after is not None else 0
        sb = para.space_before.pt if para.space_before is not None else 0
        n = measure(txt, sz, width_in)
        total += (n * sz * ls * 1.02 + sa + sb) / 72.0
    return total, maxsize


def main(path):
    prs = Presentation(path)
    sw = prs.slide_width / EMU_IN
    problems, small = [], []
    for idx, slide in enumerate(prs.slides, 1):
        boxes = []
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            txt = shp.text_frame.text.strip()
            if not txt:
                continue
            x = shp.left / EMU_IN; y = shp.top / EMU_IN
            w = shp.width / EMU_IN; h = shp.height / EMU_IN
            inner = w - (shp.text_frame.margin_left + shp.text_frame.margin_right) / EMU_IN
            need, maxsize = frame_height(shp.text_frame, max(inner, 0.4))
            bottom = y + need
            boxes.append((x, y, w, h, need, bottom, txt[:46]))
            if bottom > SAFE_BOTTOM:
                problems.append(f"  s{idx:>3} OFF-SLIDE  bottom={bottom:.2f}  «{txt[:52]}»")
            elif bottom > FOOT_Y + 0.03 and y < FOOT_Y - 0.1:
                problems.append(f"  s{idx:>3} past-footer bottom={bottom:.2f}  «{txt[:52]}»")
            if need > h + 0.30 and y + h < FOOT_Y - 0.2:
                problems.append(f"  s{idx:>3} SPILL box_h={h:.2f} need={need:.2f} «{txt[:46]}»")
            for r in shp.text_frame.paragraphs:
                for run in r.runs:
                    if run.font.size and run.font.size.pt < 18.5 and len(run.text.strip()) > 26:
                        small.append(f"  s{idx:>3} {run.font.size.pt}pt «{run.text[:48]}»")
        # pairwise vertical collision on overlapping x-ranges
        for i in range(len(boxes)):
            for j in range(len(boxes)):
                if i == j:
                    continue
                ax, ay, aw, ah, an, ab, at = boxes[i]
                bx, by, bw, bh, bn, bb, bt = boxes[j]
                if by <= ay:
                    continue
                if ax < bx + bw - 0.05 and bx < ax + aw - 0.05:
                    if ab > by + 0.06:
                        problems.append(
                            f"  s{idx:>3} COLLIDE  «{at[:30]}» bottom={ab:.2f} over «{bt[:30]}» top={by:.2f}")
    print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}" if False else f"slides checked: {idx}")
    print(f"layout problems: {len(problems)}")
    for p in problems:
        print(p)
    print(f"body runs under 18.5pt: {len(small)}")
    for p in small[:25]:
        print(p)


if __name__ == "__main__":
    main(sys.argv[1])
