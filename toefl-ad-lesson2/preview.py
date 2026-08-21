# -*- coding: utf-8 -*-
"""Rough raster preview of the deck (no LibreOffice in this container).

Draws rectangles, connectors and wrapped text straight from the pptx XML so the
design can be eyeballed. Approximate, but faithful enough to catch ugly slides.
"""
import sys, os
from pptx import Presentation
from pptx.util import Pt
from PIL import Image, ImageDraw, ImageFont

EMU = 914400.0
DPI = 110
FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
_fc = {}


def f(px, bold=False):
    key = (int(px), bold)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(FONT, max(6, int(px)))
    return _fc[key]


def rgb(c):
    try:
        return (c[0], c[1], c[2])
    except Exception:
        return None


def shape_fill(shp):
    try:
        if shp.fill.type is not None and shp.fill.type == 1:
            return tuple(shp.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def shape_line(shp):
    try:
        if shp.line.fill.type == 1:
            return tuple(shp.line.color.rgb), max(1, int(shp.line.width.pt * DPI / 72)) if shp.line.width else 1
    except Exception:
        pass
    return None, 0


def wrap(draw, text, font, maxw):
    toks, buf = [], ""
    for ch in text:
        if ord(ch) > 0x2E80:
            if buf:
                toks.append(buf); buf = ""
            toks.append(ch)
        elif ch == " ":
            buf += ch; toks.append(buf); buf = ""
        else:
            buf += ch
    if buf:
        toks.append(buf)
    out, cur = [], ""
    for t in toks:
        trial = cur + t
        if font.getlength(trial.rstrip()) > maxw and cur.strip():
            out.append(cur.rstrip()); cur = t.lstrip()
        else:
            cur = trial
    out.append(cur)
    return out


def render(path, outdir, only=None):
    prs = Presentation(path)
    W = int(prs.slide_width / EMU * DPI)
    H = int(prs.slide_height / EMU * DPI)
    os.makedirs(outdir, exist_ok=True)
    for i, slide in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        img = Image.new("RGB", (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for shp in slide.shapes:
            x = int(shp.left / EMU * DPI); y = int(shp.top / EMU * DPI)
            w = int(shp.width / EMU * DPI); h = int(shp.height / EMU * DPI)
            if shp.shape_type is not None and str(shp.shape_type).startswith("LINE"):
                lc, lw = shape_line(shp)
                d.line([x, y, x + w, y + h], fill=lc or (200, 200, 200), width=lw or 1)
                continue
            fl = shape_fill(shp)
            lc, lw = shape_line(shp)
            if fl or lc:
                d.rounded_rectangle([x, y, x + w, y + h], radius=6, fill=fl, outline=lc, width=lw or 1)
            if not shp.has_text_frame:
                continue
            tf = shp.text_frame
            ml = tf.margin_left / EMU * DPI
            cy = y + tf.margin_top / EMU * DPI
            maxw = w - ml - tf.margin_right / EMU * DPI
            for para in tf.paragraphs:
                runs = [(r.text, r.font) for r in para.runs if r.text]
                if not runs:
                    cy += 6
                    continue
                sz = max([rn[1].size.pt for rn in runs if rn[1].size] or [18])
                ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.16
                px = sz * DPI / 72.0
                full = "".join(rn[0] for rn in runs)
                col = None
                for rn in runs:
                    try:
                        col = tuple(rn[1].color.rgb); break
                    except Exception:
                        pass
                bold = any(rn[1].bold for rn in runs)
                fnt = f(px, bold)
                lines = wrap(d, full, fnt, maxw)
                for ln in lines:
                    tw = fnt.getlength(ln)
                    al = str(para.alignment or "")
                    if "CENTER" in al:
                        tx = x + ml + (maxw - tw) / 2
                    elif "RIGHT" in al:
                        tx = x + ml + maxw - tw
                    else:
                        tx = x + ml
                    d.text((tx, cy), ln, font=fnt, fill=col or (20, 20, 20))
                    cy += px * ls
                cy += (para.space_after.pt if para.space_after else 0) * DPI / 72.0
        img.save(f"{outdir}/s{i:03d}.png")
    print("rendered ->", outdir)


if __name__ == "__main__":
    only = set(int(v) for v in sys.argv[3].split(",")) if len(sys.argv) > 3 else None
    render(sys.argv[1], sys.argv[2], only)
