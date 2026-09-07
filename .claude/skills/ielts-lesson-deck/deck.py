# -*- coding: utf-8 -*-
"""雅思阅读精讲班 · 课件版式引擎。

设计令牌沿用第二节课 PPT（深色底 + 顶部蓝条 + 四色面板）。
本引擎负责"怎么摆"，内容由各课的 teach_*.py 提供，两者不耦合。

相比第三节课的初版，本版做了五处优化：
  1. 分字宽模型      —— 按字符类别估算宽度，不再把所有西文当 0.55em，减少无谓缩字
  2. normAutofit    —— 写入 PowerPoint 自带的"溢出自动缩字"，估算失手时兜底
  3. 讲师备注        —— 逐题页自动生成备注（答案 / 陷阱 / 提问脚本）
  4. 进度指示        —— 逐题页页脚显示"第 n / N 题"
  5. qa_pair()      —— 把"一题两页"的版式收进引擎，各课不再各自硬编码坐标
"""

import math

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

# ================================================================= 设计令牌
SLIDE_W, SLIDE_H = 13.333, 7.5
BG = "0B1220"          # 底色
BAR = "38BDF8"         # 顶部条
KICKER_C = "38BDF8"    # 眉标
TITLE_C = "FFFFFF"
FOOT_C = "94A3B8"
INK = "0F172A"         # 面板内文字

CJK = "微软雅黑"
LAT = "Arial"

# tone -> (面板底色, 边框色, 标签色)
PAL = {
    "blue":   ("DBEAFE", "38BDF8", "38BDF8"),
    "purple": ("F3E8FF", "A78BFA", "C4B5FD"),
    "green":  ("DCFCE7", "22C55E", "86EFAC"),
    "amber":  ("FEF3C7", "F59E0B", "FDE68A"),
    "cream":  ("FFFBEB", "F59E0B", "FDE68A"),
    "red":    ("FEE2E2", "F87171", "FCA5A5"),
}

# ============================================================ 文本宽度估算
# 初版把所有西文按 0.55em 计算，导致长英文句子被高估、字号被压得过小。
# 这里按字符类别给宽度，估算更贴近真实，可用更大的字号。
_NARROW = set("iljItf.,;:'\"!|()[]{}·`^ ")
_WIDE = set("mwMWQ@%&")


def _char_em(ch):
    o = ord(ch)
    if 0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF:      # 汉字
        return 1.0
    if 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFFEF:      # 全角标点
        return 1.0
    if o in (0x2014, 0x2015):                               # 破折号
        return 1.0
    if ch in _NARROW:
        return 0.30
    if ch in _WIDE:
        return 0.85
    if ch.isupper():
        return 0.68
    if ch.isdigit():
        return 0.55
    return 0.52


def text_width_em(s):
    """一行文字的宽度（单位 em）。"""
    return sum(_char_em(c) for c in s)


def lines_needed(paras, width_in, size_pt):
    """给定宽度与字号，估算总行数。"""
    em = size_pt / 72.0
    usable = max(width_in - 0.14, 0.4)
    total = 0
    for p in paras:
        if not p.strip():
            total += 1
            continue
        total += max(1, math.ceil(text_width_em(p) * em / usable - 1e-6))
    return total


def fit_size(paras, width_in, height_in, sizes):
    """从 sizes（由大到小）里挑第一个放得下的字号。"""
    for s in sizes:
        if lines_needed(paras, width_in, s) * (s / 72.0) * 1.26 <= height_in:
            return s
    return sizes[-1]


# 向后兼容第三节课的内部名
_text_width_em = text_width_em
_lines_needed = lines_needed


def _rgb(h):
    return RGBColor.from_string(h)


# ================================================================== 版式表
# 逐题两页的坐标集中在这里，新增课次不需要再算一遍。
LAYOUT = {
    # 普通题（判断 / 填空 / 简答 / 表格 / 配对）
    "qa": dict(stem=(0.72, 1.42, 11.90, 1.02),
               row2_y=2.72, row2_h=1.52,
               row3_y=4.86, row3_h=1.62),
    # 选择题：题干下面多一栏选项
    "qa_mcq": dict(stem=(0.72, 1.40, 11.90, 0.58),
                   opts=(0.72, 2.06, 11.90, 1.34),
                   row2_y=3.56, row2_h=1.28,
                   row3_y=5.38, row3_h=1.16),
    # 长难句页
    "sent": dict(sent=(0.72, 1.42, 11.90, 1.32),
                 cols_y=2.92, cols_h=2.12,
                 core=(0.78, 3.24), layers=(4.32, 4.42), cn=(9.04, 3.51),
                 vocab_y=5.55, vocab_h=1.00),
}

MCQ_LETTERS = "ABCDE"


class Deck:
    """一份课件。所有坐标单位为英寸。"""

    def __init__(self, footer="雅思阅读精讲 · 黄玉蕾 Renee", notes=True, autofit=True):
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(SLIDE_W * 914400))
        self.prs.slide_height = Emu(int(SLIDE_H * 914400))
        self.blank = self.prs.slide_layouts[6]
        self.footer = footer
        self.notes_on = notes
        self.autofit = autofit
        self.n = 0

    # ------------------------------------------------------------ 低层绘制
    def _box(self, slide, x, y, w, h, fill=None, line=None):
        sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                    Inches(w), Inches(h))
        sp.shadow.inherit = False
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = _rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = _rgb(line)
            sp.line.width = Pt(1)
        else:
            sp.line.fill.background()
        return sp

    def _write(self, sp, paras, size, color, bold=False, align=PP_ALIGN.LEFT,
               anchor=MSO_ANCHOR.MIDDLE, inset=0.06, latin=CJK, autofit=None):
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = Inches(inset)
        tf.margin_right = Inches(inset)
        tf.margin_top = Inches(0.03)
        tf.margin_bottom = Inches(0.03)
        # 估算失手时，让 PowerPoint 自己把字缩进框里（见 SKILL.md 的说明）
        if (self.autofit if autofit is None else autofit):
            try:
                tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            except Exception:
                pass
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt
            f = r.font
            f.size = Pt(size)
            f.bold = bold
            f.color.rgb = _rgb(color)
            f.name = latin
            # 东亚字体必须显式指定，否则 PowerPoint 会回退成默认字体
            rPr = r._r.get_or_add_rPr()
            for tag in ("ea", "cs"):
                el = rPr.makeelement(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag,
                    {"typeface": latin})
                rPr.append(el)
        return sp

    def _label(self, slide, x, y, text, color, w=3.4):
        sp = self._box(slide, x, y, w, 0.30)
        self._write(sp, [text], 18.5, color, bold=True, inset=0.0, autofit=False)
        return sp

    def note(self, slide, lines):
        """写讲师备注。lines 可为字符串或字符串列表。"""
        if not self.notes_on:
            return
        if isinstance(lines, str):
            lines = [lines]
        txt = "\n".join(x for x in lines if x)
        if txt:
            slide.notes_slide.notes_text_frame.text = txt

    # -------------------------------------------------------------- 面板
    def panel(self, slide, x, y, w, h, paras, tone, size=None, bold=False,
              sizes=(19, 18, 17, 16, 15, 14, 13, 12, 11), anchor=None):
        fill, line, _ = PAL[tone]
        sp = self._box(slide, x, y, w, h, fill, line)
        s = size or fit_size(paras, w, h, list(sizes))
        if anchor is None:
            used = lines_needed(paras, w, s) * (s / 72.0) * 1.26
            anchor = MSO_ANCHOR.MIDDLE if used < h * 0.72 else MSO_ANCHOR.TOP
        self._write(sp, paras, s, INK, bold=bold, anchor=anchor)
        return sp

    def labelled(self, slide, x, y, w, h, label, paras, tone, bold=False,
                 sizes=(19, 18, 17, 16, 15, 14, 13, 12, 11), lw=None, anchor=None):
        _, _, lab = PAL[tone]
        self._label(slide, x, y, label, lab, w=lw or w)
        return self.panel(slide, x, y + 0.34, w, h, paras, tone, bold=bold,
                          sizes=sizes, anchor=anchor)

    # -------------------------------------------------------------- 骨架
    def slide(self, kicker, title, footer=None, progress=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        self._box(s, 0, 0, SLIDE_W, SLIDE_H, BG, BG)
        self._box(s, 0, 0, SLIDE_W, 0.13, BAR, BAR)

        k = self._box(s, 0.62, 0.34, 8.6, 0.32)
        self._write(k, [kicker], 18.5, KICKER_C, bold=True, inset=0.0,
                    latin=LAT, autofit=False)

        t = self._box(s, 0.62, 0.74, 12.05, 0.66)
        self._write(t, [title], fit_size([title], 12.05, 0.66, [30, 27, 24, 22, 20]),
                    TITLE_C, bold=True, inset=0.0)

        foot = footer or self.footer
        if progress:
            foot = f"{foot}　·　{progress}"
        f = self._box(s, 0.62, 7.06, 10.9, 0.30)
        self._write(f, [foot], 15, FOOT_C, inset=0.0, autofit=False)

        pn = self._box(s, 12.15, 7.06, 0.55, 0.30)
        self._write(pn, [str(self.n)], 15, FOOT_C, align=PP_ALIGN.RIGHT,
                    inset=0.0, latin=LAT, autofit=False)
        return s

    def divider(self, kicker, title, lines=None, tone="cream"):
        """章节分隔页：长课件里用来切分模块。"""
        s = self.slide(kicker, title)
        self.panel(s, 0.80, 2.55, 11.75, 2.10, lines or [title], tone,
                   bold=True, sizes=(30, 27, 24, 22, 20))
        return s

    def three_col(self, kicker, title, cols, footer=None):
        """三栏卡片（文章地图 / 目标 / 小结通用）。cols = [(标题, [条目...]), ×3]"""
        s = self.slide(kicker, title, footer)
        for (head, items), x, tone in zip(cols, (0.72, 4.74, 8.76),
                                          ("blue", "purple", "green")):
            self._label(s, x, 1.62, head, PAL[tone][2])
            self.panel(s, x, 2.00, 3.75, 3.95, [f"· {i}" for i in items], tone,
                       bold=True, sizes=(22, 20, 18, 17, 16, 15, 14))
        return s

    def bullets(self, kicker, title, blocks, footer=None):
        """整幅若干长条（方法页通用）。blocks = [(文本或文本列表, tone), ...]"""
        s = self.slide(kicker, title, footer)
        top, bottom, gap = 1.55, 6.85, 0.22
        h = (bottom - top - gap * (len(blocks) - 1)) / len(blocks)
        y = top
        for text, tone in blocks:
            paras = text if isinstance(text, list) else [text]
            self.panel(s, 0.80, y, 11.75, h, paras, tone, bold=True,
                       sizes=(23, 21, 20, 19, 18, 17, 16, 15))
            y += h + gap
        return s

    # ================================================== 逐题两页（核心版式）
    def qa_pair(self, q, footer=None, progress=None):
        """一题两页：逐题标准精讲 + 长难句拆解/随题词汇。

        q 的字段见 reference/content-schema.md。返回 (页1, 页2)。
        """
        s1 = self._qa_answer_slide(q, footer, progress)
        s2 = self._qa_sentence_slide(q, footer, progress)
        return s1, s2

    def _qa_answer_slide(self, q, footer, progress):
        is_mcq = q.get("kind") in ("MCQ",)
        L = LAYOUT["qa_mcq"] if is_mcq else LAYOUT["qa"]
        s = self.slide("逐题标准精讲", q["title"], footer, progress)

        x, y, w, h = L["stem"]
        self.panel(s, x, y, w, h, [q["stem"]], "cream", bold=True,
                   sizes=(24, 22, 20, 18, 17, 16) if not is_mcq else (23, 21, 19, 18, 17))
        if is_mcq:
            ox, oy, ow, oh = L["opts"]
            self.panel(s, ox, oy, ow, oh,
                       [f"{MCQ_LETTERS[i]}  {o}" for i, o in enumerate(q["options"])],
                       "amber", sizes=(18, 17, 16, 15, 14))

        y2, h2 = L["row2_y"], L["row2_h"]
        y3, h3 = L["row3_y"], L["row3_h"]
        self.labelled(s, 0.78, y2, 3.55, h2, "① 入口", q["entry"], "purple", lw=2.4)
        self.labelled(s, 4.70, y2, 7.85, h2, "② 原文证据", [q["ev"]], "blue", lw=2.4)
        self.labelled(s, 0.78, y3, 7.30, h3, "③ 同义替换 / 判断逻辑", q["logic"],
                      "green", lw=3.4)
        self.labelled(s, 8.45, y3, 4.10, h3, "④ 答案",
                      [f"Answer: {q['ans']}"] + list(q.get("ansnote", [])),
                      "amber", bold=True, sizes=(21, 20, 19, 18, 17, 16), lw=2.4)

        self.note(s, [
            f"【Q{q['n']}｜{q.get('kind', '')}】答案：{q['ans']}",
            f"定位：{q['ev'].split(':')[0]}",
            "提问脚本：先问学生「题干里哪个词限定了范围？」再问「原文哪一句能独立支持它？」",
            "错因归类：定位错 / 替换错 / 范围与方向错——让学生自己说属于哪一类。",
        ])
        return s

    def _qa_sentence_slide(self, q, footer, progress):
        L = LAYOUT["sent"]
        s = self.slide("长难句拆解 + 随题词汇", q["stitle"], footer, progress)

        x, y, w, h = L["sent"]
        self.panel(s, x, y, w, h, [q["sent"]], "cream", bold=True,
                   sizes=(20.5, 19, 18, 17, 16, 15, 14))

        cy, ch = L["cols_y"], L["cols_h"]
        cx, cw = L["core"]
        self.labelled(s, cx, cy, cw, ch, "核心主干", [q["core"]], "red",
                      bold=True, sizes=(20, 19, 18, 17, 16, 15, 14), lw=2.0)
        lx, lw_ = L["layers"]
        self.labelled(s, lx, cy, lw_, ch, "层层拆解", [f"· {x}" for x in q["layers"]],
                      "blue", sizes=(18.8, 18, 17, 16, 15, 14, 13), lw=2.0)
        nx, nw = L["cn"]
        self.labelled(s, nx, cy, nw, ch, "课堂表达", [q["cn"]], "green",
                      sizes=(19.2, 18, 17, 16, 15, 14, 13), lw=2.0)

        self.labelled(s, 0.78, L["vocab_y"], 11.77, L["vocab_h"],
                      "随题词汇 / 同义替换", q["vocab"], "amber", bold=True,
                      sizes=(18.5, 17.5, 16.5, 15.5, 14.5), lw=3.4)

        self.note(s, [
            f"【Q{q['n']} 长难句】主干：{q['core']}",
            "课堂动作：先让学生圈出有限谓语，再盖住插入语读主干，最后才看翻译。",
            "板书：主干写一行，从句用括号括起来，比较/否定单独标箭头。",
        ])
        return s

    def passage_block(self, mod, no, title_en, title_cn, footer,
                      hooks=None, label=None, head=None):
        """一篇文章：文章地图 → 题型策略 → 逐题两页 → 小结。"""
        hooks = hooks or {}
        name = label or f"Passage {no}"
        total = len(mod.Q)
        self.three_col("文章地图", head or f"PASSAGE {no} · {title_en}", mod.MAP)
        self.bullets("题型策略", f"{name} 题型策略",
                     [(t[0], tone) for t, tone in
                      zip(mod.STRATEGY, ["cream", "blue", "purple", "green"])])
        for i, q in enumerate(mod.Q, 1):
            if q["n"] in hooks:
                hooks[q["n"]](self, footer)
            self.qa_pair(q, footer, progress=f"第 {i} / {total} 题")
        self.three_col("小结", f"{name} 小结：{title_cn}", mod.SUMMARY)

    # ---------------------------------------------------------- 自定义图形
    def wheel(self, slide, cx, cy, r, style, color=INK, spokes=6):
        """轮辐示意图（剑桥 4 Test 1 图形配对题专用）。

        style: 'extend' 伸出轮圈 / 'dashed' 虚线 / 'curved' 弧线。
        其他课如需自定义图形，照此模式另写一个方法即可。
        """
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.enum.dml import MSO_LINE_DASH_STYLE

        ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                                      Inches(2 * r), Inches(2 * r))
        ring.shadow.inherit = False
        ring.fill.background()
        ring.line.color.rgb = _rgb(color)
        ring.line.width = Pt(1.75)

        for i in range(spokes):
            a = math.pi * 2 * i / spokes
            if style == "curved":
                pts = []
                for j in range(11):
                    t = j / 10
                    ang = a + 0.62 * t
                    pts.append((cx + r * t * math.cos(ang), cy + r * t * math.sin(ang)))
                ff = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
                ff.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]],
                                     close=False)
                sp = ff.convert_to_shape()
                sp.shadow.inherit = False
                sp.fill.background()
                sp.line.color.rgb = _rgb(color)
                sp.line.width = Pt(1.75)
            else:
                ext = 1.34 if style == "extend" else 1.0
                cn = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(cy),
                    Inches(cx + r * ext * math.cos(a)),
                    Inches(cy + r * ext * math.sin(a)))
                cn.shadow.inherit = False
                cn.line.color.rgb = _rgb(color)
                cn.line.width = Pt(1.75)
                if style == "dashed":
                    cn.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def save(self, path):
        self.prs.save(path)
        return path
